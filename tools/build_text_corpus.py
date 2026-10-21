"""Build local CSV resources from pinned assembly and translator workbooks.

Development data construction only (Python + openpyxl); never reads testpic,
OCR output, ROMs or saves, and never compiles/executes upstream game code.
Sources live in ignored output/corpus-sources/<id>. See corpus/README.md.
"""
import collections
import csv
import json
import pathlib
import re
import subprocess

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parents[1]
SOURCES = ROOT / 'output/corpus-sources'
DEST = ROOT / 'src/main/resources/corpus'
TEXT_OPS = {'text', 'text_start', 'line', 'next', 'cont', 'para', '_cnt', 'scrl',
            'text_ram', 'text_decimal', 'text_bcd', 'text_dots', 'text_end',
            'done', 'prompt', 'text_low', 'text_scroll', 'text_promptbutton',
            'text_waitbutton', 'text_pause', 'text_move'}
JP_FIXED = {'#': 'ポケモン', '<PC>': 'パソコン', '<TM>': 'わざマシン',
            '<TRAINER>': 'トレーナー', '<ROCKET>': 'ロケットだん', '<⋯>': '⋯⋯'}
ZH_FIXED = {'#MON': '宝可梦', '#': '宝可梦', '&': '训练家', '+': '火箭队'}
BOOKS = ['buildingsA', 'buildingsB', 'core', 'indoor', 'outdoor', 'routes', 'ratings']


def write_csv(path, rows, fields):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=fields, lineterminator='\n', extrasaction='ignore')
        writer.writeheader()
        writer.writerows(rows)


def packed(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'))


def literal(value):
    # RGBDS Unicode escapes are source notation, not visible backslashes.
    return re.sub(r'\\u([0-9a-fA-F]{4})', lambda m: chr(int(m[1], 16)), value)


def english_key(value):
    # Offline bridge only; this is never an OCR lookup normalization.
    value = value.replace('#MON', 'POKEMON').replace('<PKMN>', 'POKEMON')
    return re.sub(r'\s+', '', value.replace('é', 'e').replace('@', '')).lower()


def label_key(value):
    return value.strip(':_')


def extract_assembly(repo, game):
    records, edges = [], collections.defaultdict(list)
    flags = {'_'+game.upper(), '_'+game.upper()+'_VC', '_REV3' if game == 'yellow' else '_REV1'}
    for folder in ['scripts', 'engine', 'home', 'data', 'text']:
        for path in sorted((SOURCES/repo/folder).rglob('*.asm')):
            global_label, label, current = '', '', None
            conditions = []
            for ln, raw in enumerate(path.read_text(encoding='utf-8').splitlines(), 1):
                stripped = raw.strip()
                if stripped.startswith('IF '):
                    # The pinned text-bearing files use DEF build selectors.
                    terms = stripped[3:].split('||')
                    conditions.append(any((not bool(set(re.findall(r'DEF\s*\((\w+)\)', term)) & flags))
                                          if term.strip().startswith('!') else
                                          bool(set(re.findall(r'DEF\s*\((\w+)\)', term)) & flags)
                                          for term in terms))
                    continue
                if stripped == 'ELSE':
                    conditions[-1] = not conditions[-1]
                    continue
                if stripped == 'ENDC':
                    conditions.pop()
                    continue
                if conditions and not all(conditions):
                    continue
                match = re.match(r'^([.\w]+)::?', raw) or re.fullmatch(r'(\.\w+)\s*(?:;.*)?', raw)
                if match:
                    if current:
                        records.append(current)
                        current = None
                    label = match[1]
                    if label.startswith('.'):
                        label = global_label + label
                    else:
                        global_label = label
                far = re.search(r'\btext_far\s+(\w+)', raw)
                if far:
                    edges[label].append(far[1])
                match = re.match(r'\s+(\w+)\b\s*(.*)', raw)
                if not match:
                    continue
                op, arg = match[1], match[2].strip()
                if op == 'text_asm':
                    if current:
                        current['termination'] = 'asm'
                        records.append(current)
                        current = None
                elif op in TEXT_OPS or op.startswith('sound_'):
                    if current is None:
                        if op not in {'text', 'text_start', 'text_ram', 'text_decimal', 'text_bcd'}:
                            continue
                        current = dict(repo=repo, file=path.relative_to(SOURCES/repo).as_posix(),
                                       label=label, line=ln, ops=[], termination='fallthrough')
                    current['ops'].append([op, arg, ln])
                    if op in {'done', 'prompt', 'text_end'}:
                        current['termination'] = op
                        records.append(current)
                        current = None
                elif raw.strip() and not raw.lstrip().startswith(';'):
                    if current:
                        records.append(current)
                        current = None
            if current:
                records.append(current)
    return records, edges


def quoted_text(ops):
    return ''.join(literal(''.join(re.findall(r'"([^"\n]*)"', arg)))
                   for op, arg, _ in ops if op in {'text', 'line', 'next', 'cont', 'para', '_cnt', 'scrl'})


def extract_workbooks(repo, version, game):
    blocks, replacements = [], []
    for book in BOOKS:
        path = SOURCES/repo/'src/xlsx'/f'{book}.xlsx'
        if not path.exists():
            continue  # Red/Green/Blue has no ratings.xlsx; Yellow does.
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in wb:
            block = None
            source_file = sheet.cell(1, 1).value
            conditions = []
            for rowno, values in enumerate(sheet.iter_rows(values_only=True), 1):
                row = list(values) + [None]*16
                inst = str(row[5] or '').strip()
                if inst.startswith('IF '):
                    flags = set(re.findall(r'DEF\((\w+)\)', inst))
                    enabled = {'_'+game.upper(), '_'+game.upper()+'_VC'}
                    conditions.append(bool(flags & enabled))
                    continue
                if inst == 'ELSE':
                    conditions[-1] = not conditions[-1]
                    continue
                if inst == 'ENDC':
                    conditions.pop()
                    continue
                if conditions and not all(conditions):
                    continue
                if row[4] == 'end':
                    break
                if row[3] and version not in str(row[3]):
                    continue  # Same D-column build selector as _importText.py mode 5.
                if isinstance(row[4], str) and row[4].endswith(':'):
                    block = dict(repo=repo, book=book+'.xlsx', sheet=sheet.title,
                                 row=rowno, label=row[4], file=source_file,
                                 ops=[], english=[], cells=[], reference=[])
                    blocks.append(block)
                if block is None:
                    continue
                if row[5] is not None:
                    op, arg = str(row[5]), '' if row[6] is None else str(row[6])
                    block['ops'].append([op, arg, rowno])
                    block['cells'].append([f'F{rowno}', row[5], f'G{rowno}', row[6]])
                if row[1] is not None:
                    block['english'].append([str(row[1]), '' if row[2] is None else str(row[2]), rowno])
                # I is the translator's Japanese reference in RGB; Yellow's
                # automatically matched Gen1:JPN column is deliberately not used.
                if repo == 'pokeredCHS-rgb' and row[8] is not None:
                    block['reference'].append(str(row[8]))
        wb.close()
    path = SOURCES/repo/'src/xlsx/data.xlsx'
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in wb:
        file, ending, selected = '', '', True
        for rn, values in enumerate(sheet.iter_rows(values_only=True), 1):
            row = list(values) + [None]*6
            if row[0] == 'end':
                break
            if row[0] is not None:
                file, ending = str(row[0]), str(row[1] or '')
                selected = not row[2] or version in str(row[2])
                continue
            if selected and row[1] is not None:
                replacements.append(dict(repo=repo, sheet=sheet.title, row=rn, file=file,
                                         original=str(row[1]), chinese=str(row[2] or ''),
                                         ending=ending, replacement=row[3], context=row[4]))
    wb.close()
    return blocks, replacements


def render(ops, language, quoted=True):
    """Render the finite text-command subset; preserve layout events separately."""
    text, events, variables, unsupported = '', [], [], []
    counts, slots = collections.Counter(), {}
    for op, arg, source_line in ops:
        start = len(text)
        if op.startswith(';') or op == '\\':
            continue  # Source comment / assembler line continuation, not glyphs.
        if op == 'text $4c,':
            events.append(dict(offset=len(text),op='<SCROLL>',source_line=source_line))
            text += '\n'
            op = 'text'
        if op in {'line', 'next', 'cont', '_cnt', 'scrl', 'para'}:
            text += '\n\n' if op == 'para' else '\n'
            events.append(dict(offset=start, op=op, source_line=source_line))
        if op in {'text', 'line', 'next', 'cont', '_cnt', 'scrl', 'para'}:
            value = literal(''.join(re.findall(r'"([^"\n]*)"', arg))) if quoted else literal(arg)
            # @ ends a PlaceString call; following RAM/START commands continue.
            value = value.replace('@', '')
            for token, expansion in (JP_FIXED if language == 'ja' else ZH_FIXED).items():
                value = value.replace(token, expansion)
            for token in ['<_CONT>', '<CONT>', '<SCROLL>', '<LINE>', '<NEXT>', '<PARA>']:
                if token in value:
                    events.append(dict(offset=len(text)+value.index(token), op=token, source_line=source_line))
                    value = value.replace(token, '\n\n' if token == '<PARA>' else '\n')
            for token in ['PLAYER', 'RIVAL', 'USER', 'TARGET']:
                value = value.replace('<'+token+'>', '{{'+token+'}}')
            text += value
        elif op in {'text_ram', 'text_decimal', 'text_bcd'}:
            address = arg.split(',')[0].strip()
            kind = 'RAM' if op == 'text_ram' else 'NUMBER'
            if (kind, address) not in slots:
                counts[kind] += 1
                slots[kind, address] = f'{kind}{counts[kind]}'
            slot = slots[kind, address]
            variables.append(dict(slot=slot, kind=kind, address=address, format=arg,
                                  offset=len(text), source_line=source_line))
            text += '{{'+slot+'}}'
        elif op == 'text_dots':
            text += '⋯' * int(arg.split(';')[0].strip())
        elif op in {'text_scroll', 'text_low'}:
            events.append(dict(offset=len(text), op=op, source_line=source_line))
            text += '\n'
        elif op not in TEXT_OPS and not op.startswith('sound_'):
            # Assembly/conditional controls are retained and never silently
            # certified as an ordinary renderable text template.
            unsupported.append(op)
    for m in re.finditer(r'\{\{(PLAYER|RIVAL|USER|TARGET)\}\}', text):
        variables.append(dict(slot=m[1], kind='NAME', address='', format='', offset=m.start()))
    if re.search(r'<[^>]+>', text):
        unsupported += re.findall(r'<[^>]+>', text)
    return text, events, variables, sorted(set(unsupported))


def direct_key(ops):
    # Workbook English cells are unquoted. Text content only, not RAM addresses.
    return english_key(''.join(arg for op, arg, _ in ops
                              if op in {'text', 'line', 'cont', 'para', 'next'}))


def constants(repo, filename):
    values, value = {}, 0
    for line in (SOURCES/repo/filename).read_text(encoding='utf-8').splitlines():
        line = line.split(';')[0].strip()
        if filename.endswith('item_constants.asm') and line.startswith('DEF NUM_ITEMS '):
            break
        match = re.fullmatch(r'const_def(?:\s+(\$[0-9a-fA-F]+|\d+))?', line)
        if match:
            value = int((match[1] or '0').replace('$', '0x'), 0)
        elif line.startswith('const_skip'):
            value += int(line.removeprefix('const_skip').strip() or '1')
        elif match := re.fullmatch(r'(?:const|trainer_const)\s+(\w+)', line):
            values[match[1]] = value
            value += 1
    return values


def name_table(repo, filename):
    table = []
    for rn, line in enumerate((SOURCES/repo/filename).read_text(encoding='utf-8').splitlines(), 1):
        match = re.match(r'\s*(?:li|dname|db)\s+"([^"\n]*)"', line)
        if match:
            table.append((literal(match[1]).rstrip('@'), rn))
    return table


def add_words(jp, zh, games, replacements, all_records, all_fragments):
    for kind, sheet, filename, constfile in [
        ('pokemon', 'PM', 'data/pokemon/names.asm', 'constants/pokemon_constants.asm'),
        ('move', 'Moves', 'data/moves/names.asm', 'constants/move_constants.asm'),
        ('item', 'Items', 'data/items/names.asm', 'constants/item_constants.asm'),
        ('trainer', 'Trainers', 'data/trainers/names.asm', 'constants/trainer_constants.asm')]:
        jt, zt = name_table(jp, filename), name_table(zh+'/src', filename)
        jc, zc = constants(jp, constfile), constants(zh+'/src', constfile)
        translations = collections.defaultdict(list)
        for p in replacements:
            if p['sheet'] == sheet:
                translations[p['original'].strip('"').rstrip('@')].append(p)
        for symbol, number in jc.items():
            if kind == 'pokemon' and symbol in {'FOSSIL_KABUTOPS','FOSSIL_AERODACTYL','MON_GHOST'}:
                continue  # Battle graphic IDs are not ordinary species names.
            zsymbol = {'ITEM_2C':'UNUSED_ITEM','ITEM_32':'PP_UP_2'}.get(symbol,symbol)
            if number == 0 or number > len(jt) or zsymbol not in zc:
                continue
            zn = zc[zsymbol]
            if not 0 < zn <= len(zt):
                continue
            english = zt[zn-1][0]
            options = [p for p in translations[english] if p['file'] == filename]
            if not options:
                continue
            if len({p['chinese'] for p in options}) != 1:
                raise ValueError(f'Conflicting name translations: {jp} {symbol}')
            p = options[0]
            ja, ja_line = jt[number-1]
            chinese_source = p['chinese']
            chinese = literal(chinese_source).replace('@', '')
            rid = f'{jp}:{kind}:{symbol}'
            all_records.append(dict(id=rid, games=games, kind=kind, japanese_repo=jp,
                japanese_file=filename,japanese_label=symbol,japanese_line=ja_line,
                japanese=ja,japanese_commands=packed([['name',jt[number-1][0],ja_line]]),
                japanese_events='[]',japanese_termination='name',chinese_repo=zh,
                chinese_source=f"src/xlsx/data.xlsx#{sheet}!C{p['row']}",
                chinese_label=symbol,chinese=chinese,chinese_cells=packed([[f"C{p['row']}",chinese_source]]),
                chinese_events='[]',variables='[]',method='symbolic_id_and_workbook_literal',
                evidence=packed(dict(symbol=symbol,japanese_id=number,chinese_id=zn,
                                     original_english=english,constant_file=constfile,chinese_symbol=zsymbol)),
                mapping_status='aligned',segment_status='whole_visible_unit',
                variable_status='none',render_status='parsed'))
            all_fragments.append(dict(record_id=rid,ordinal=1,japanese=ja,previous_ordinal='',
                chinese=chinese,chinese_start=0,chinese_end=len(chinese),
                status='whole_visible_unit',source='symbolic_item_move_species'))
    # Name suffixes are produced by GetMachineName, not a table of 55 strings.
    for family, ja_prefix, total in [('TM','わざマシン',50),('HM','ひでんマシン',5)]:
        for n in range(1,total+1):
            rid=f'{jp}:item:{family}{n:02d}'
            ja, chinese = ja_prefix+str(n).zfill(2).translate(str.maketrans('0123456789','０１２３４５６７８９')),f'{family}{n:02d}'
            all_records.append(dict(id=rid,games=games,kind='item',japanese_repo=jp,
                japanese_file='home/names.asm',japanese_label=family+str(n).zfill(2),japanese_line=56,
                japanese=ja,japanese_commands='[]',japanese_events='[]',japanese_termination='name',
                chinese_repo=zh,chinese_source='src/home/names.asm#GetMachineName',chinese_label=family,
                chinese=chinese,chinese_cells=packed(dict(prefix=family,number=f'{n:02d}')),
                chinese_events='[]',variables='[]',method='source_machine_name_algorithm',
                evidence='GetMachineName copies TM/HM and writes two decimal digits',
                mapping_status='aligned',segment_status='whole_visible_unit',variable_status='none',render_status='parsed'))
            all_fragments.append(dict(record_id=rid,ordinal=1,japanese=ja,previous_ordinal='',
                chinese=chinese,chinese_start=0,chinese_end=len(chinese),status='whole_visible_unit',source='GetMachineName'))


def scalar_strings(repo, filename):
    values, global_label, label = {}, '', ''
    for ln, line in enumerate((SOURCES/repo/filename).read_text(encoding='utf-8').splitlines(),1):
        m=re.match(r'^([.\w]+):+',line)
        if m:
            label=m[1]
            if label.startswith('.'):label=global_label+label
            else:global_label=label
        m=re.search(r'\bdb\s+"([^"\n]*)"',line)
        if m:values[label]=(literal(m[1]),ln)
    return values


def add_fixed_words(jp, zh, games, replacements, all_records, all_fragments):
    template=next(r for r in all_records if r['japanese_repo']==jp and r['kind']=='item')

    def append(label, ja, p, filename, source_line, chinese=None):
        z=literal(p['chinese']).replace('@','') if chinese is None else chinese
        for a,b in ZH_FIXED.items():z=z.replace(a,b)
        row=dict(template)
        rid=f'{jp}:fixed:{label}'
        row.update(id=rid,kind='fixed',japanese_file=filename,japanese_label=label,japanese_line=source_line,
                   japanese=ja.replace('@',''),japanese_commands=packed([['db',ja,source_line]]),
                   chinese_source=f"src/xlsx/data.xlsx#{p['sheet']}!C{p['row']}",chinese_label=label,
                   chinese=z,chinese_cells=packed([[f"C{p['row']}",p['chinese']]]),
                   method='source_table_role_and_original_literal',evidence=packed(dict(original_english=p['original'],source_file=p['file'])))
        all_records.append(row)
        all_fragments.append(dict(record_id=rid,ordinal=1,japanese=row['japanese'],previous_ordinal='',
            chinese=z,chinese_start=0,chinese_end=len(z),status='whole_visible_unit',source='fixed_layout_or_variable_dictionary'))

    for filename in ['data/types/names.asm','data/trainers/name_pointers.asm','home/list_menu.asm','engine/battle/core.asm']:
        jstrings,zstrings=scalar_strings(jp,filename),scalar_strings(zh+'/src',filename)
        for label,(ja,ln) in jstrings.items():
            if filename.endswith('core.asm') and label not in {'DisabledText','TypeText'}:continue
            if filename=='home/list_menu.asm' and label!='ListMenuCancelText':continue
            if label not in zstrings:continue
            en=zstrings[label][0]
            options=[p for p in replacements if p['file']==filename and p['original']=='"'+en+'"' and not p['replacement']]
            if options:
                p=options[0]
                if not p['chinese'] and filename=='data/trainers/name_pointers.asm':
                    # The importer redirects these six abbreviated names to
                    # wTrainerName, so the blank C cell is not visible text.
                    symbol={'LassName':'LASS','JrTrainerMName':'JR_TRAINER_M','PokemaniacName':'POKEMANIAC',
                            'SuperNerdName':'SUPER_NERD','BeautyName':'BEAUTY','ChiefName':'CHIEF'}[label.split('.')[-1]]
                    full=next(r for r in all_records if r['id']==f'{jp}:trainer:{symbol}')
                    cell,value=json.loads(full['chinese_cells'])[0]
                    p=dict(p,chinese=value,row=int(cell[1:]),file='data/trainers/names.asm')
                append(label,ja,p,filename,ln)
    filename='data/battle/stat_mod_names.asm'
    for (ja,ln),(en,_) in zip(name_table(jp,filename),name_table(zh+'/src',filename)):
        p=next(p for p in replacements if p['file']==filename and p['original']=='"'+en+'"')
        append('StatMod.'+en,ja,p,filename,ln)
    filename='data/text_boxes.asm'
    first=next(p for p in replacements if p['file']==filename and p['original']=='"FIGHT <PK><MN>"')
    second=next(p for p in replacements if p['file']==filename and p['original']=='"ITEM  RUN@"')
    # The selected RGB/YEJP build puts Fight/Item on the top row and
    # Pokemon/Run on the bottom row. Each Chinese word is a source substring.
    for ja,word,p in [('たたかう','战斗',first),('どうぐ','道具',first),('ポケモン','宝可梦',second),('にげる','逃走',second)]:
        append('BattleMenu.'+ja,ja,p,filename,0,word)
    bait=next(p for p in replacements if p['file']==filename and p['original']=='"BALL×       BAIT"')
    rock=next(p for p in replacements if p['file']==filename and p['original']=='"THROW ROCK  RUN@"')
    append('SafariMenu.Bait','エサをなげる',bait,filename,0,'诱饵')
    append('SafariMenu.Rock','いしをなげる',rock,filename,0,'石块')


def add_compositions(jp, all_records, all_fragments):
    """The game's small, explicit battle-text continuations, not an ASM VM."""
    bylabel = {r['japanese_label']: r for r in all_records if r['japanese_repo'] == jp}
    consumers = collections.defaultdict(list)

    def append(name, ja_parts, zh_parts, detail):
        js, zs = [bylabel[x] for x in ja_parts], [bylabel[x] for x in zh_parts]
        jops = [op for r in js for op in json.loads(r['japanese_commands'])]
        ja, events, vs, unsupported = render(jops, 'ja')
        chinese = ''.join(r['chinese'] for r in zs)
        row = dict(js[0])
        rid = f'{jp}:composed:{name}'
        row.update(id=rid,kind='dialogue_template',japanese_label=name,japanese=ja,
                   japanese_commands=packed(jops),japanese_events=packed(events),japanese_termination='done',
                   chinese=chinese,chinese_label=';'.join(r['chinese_label'] for r in zs),
                   chinese_source=packed([r['chinese_source'] for r in zs]),
                   chinese_cells=packed([json.loads(r['chinese_cells']) for r in zs if r['chinese_cells']]),
                   chinese_events=packed([json.loads(r['chinese_events']) for r in zs if r['chinese_events']]),
                   variables=packed(vs),method='model_reviewed_source_control_flow',
                   evidence=packed(dict(japanese_components=[r['id'] for r in js],
                                        chinese_components=[r['id'] for r in zs],detail=detail,reviewer='model')),
                   mapping_status='aligned',variable_status='model_semantic',render_status='unsupported' if unsupported else 'parsed',
                   segment_status='whole_visible_unit' if len(ja.splitlines())<=2 and '\n\n' not in ja else 'unavailable')
        all_records.append(row)
        for r in js:
            consumers[r['id']].append(rid)
        visible=[]
        seq=0
        for line in ja.split('\n'):
            if not line:
                visible=[]
                continue
            visible=(visible+[line])[-2:]
            seq+=1
            all_fragments.append(dict(record_id=rid,ordinal=seq,japanese='\n'.join(visible),
                previous_ordinal=seq-1 if seq>1 else '',chinese='',chinese_start='',chinese_end='',
                status='unavailable',source='explicit_control_flow'))
        if row['segment_status']=='whole_visible_unit':
            all_fragments[-1].update(chinese=chinese,chinese_start=0,chinese_end=len(chinese),status='whole_visible_unit')

    grammar = (SOURCES/jp/'data/moves/grammar.asm').read_text(encoding='utf-8')
    grammar_sets=[[]]
    for line in grammar.splitlines():
        m=re.match(r'\s*db\s+(\w+)\s*(?:;.*)?$',line)
        if m:
            if m[1]=='0':grammar_sets.append([])
            else:grammar_sets[-1].append(m[1])
    for n in range(5):
        end = ('SentenceEnd' if jp=='pokeyellow-jp' else 'EndUsedMove')+str(n+1)+'Text'
        for disobedient in [False,True]:
            parts=['UsedMoveText','UserWaText' if disobedient or n<3 else 'UserNoText']
            zparts=['UsedMoveText']
            if disobedient:
                parts.append('UsedMoveText_CheckObedience.UsedInsteadText')
                zparts.append('UsedMoveText_CheckObedience.UsedInsteadText')
            parts+=['MoveNameText',end]
            zparts+=['MoveNameText']  # Chinese verb and ! already belong here.
            append(f'UsedMove.grammar{n}.'+('disobedient' if disobedient else 'normal'),parts,zparts,
                   dict(grammar=n,move_symbols=grammar_sets[n],default_set=n==4,
                        rule='GetMoveGrammar plus actual sentence-ending pointer table; not comments'))
    for prefix in ['GoText','DoItText','GetmText','EnemysWeakText']:
        append('SendOut.'+prefix,[prefix,'PlayerMon1Text'],[prefix,'PlayerMon1Text'],'PrintPlayerMon1Text')
    for middle in [[],['EnoughText'],['OKExclamationText'],['GoodText']]:
        parts=['PlayerMon2Text']+middle+['ComeBackText']
        append('Retreat.'+(middle[0] if middle else 'direct'),parts,parts,'RetreatMon/PrintComeBackText')
    for middle in [[],['WithExpAllText'],['BoostedText']]:
        parts=['GainedText']+middle+['ExpPointsText']
        append('Experience.'+(middle[0] if middle else 'normal'),parts,parts,'GainedText chooses its explicit continuation')
    for prefix, modifier, end in [('MonsStatsRoseText','GreatlyRoseText','RoseText'),('MonsStatsFellText','GreatlyFellText','FellText')]:
        for large in [False,True]:
            parts=[prefix]+([modifier] if large else [])+[end]
            append('StatChange.'+prefix+('.large' if large else '.normal'),parts,parts,'StatModifier effect and source fallthrough')
    for tail in ['MadeWhirlwindText','TookInSunlightText','LoweredItsHeadText','SkyAttackGlowingText','FlewUpHighText','DugAHoleText']:
        append('Charge.'+tail,['ChargeMoveEffectText',tail],[tail],'Chinese tail already includes USER; do not print the name twice')
    for row in all_records:
        if row['id'] in consumers:
            row['kind']='component'
            row['mapping_status']='component'
            row['segment_status']='use_composition'
            row['evidence']=packed(dict(alignment=row['evidence'],consumed_by=consumers[row['id']]))


def build():
    locks = list(csv.DictReader((ROOT/'tools/text_sources.csv').open(encoding='utf-8')))
    source_rows = []
    for lock in locks:
        checkout = SOURCES/lock['id']
        head = subprocess.check_output(['git', '-C', str(checkout), 'rev-parse', 'HEAD'], text=True).strip()
        if head != lock['commit']:
            raise ValueError(f"Pinned source differs: {lock['id']} {head}")
        source_rows.append(dict(lock, license_status='redistribution_not_confirmed'))
    overrides = list(csv.DictReader((ROOT/'tools/text_alignment.csv').open(encoding='utf-8')))
    bindings = {row['japanese_label']: row for row in csv.DictReader((ROOT/'tools/text_variable_bindings.csv').open(encoding='utf-8'))}
    all_records, all_fragments, all_variables, issues = [], [], [], []
    stats = {}
    for jp, en, zh, games, ver in [
        ('pokegreen', 'pokered', 'pokeredCHS-rgb', 'red;green', 'RGB'),
        ('pokeblue', 'pokered', 'pokeredCHS-rgb', 'blue', 'RGB'),
        ('pokeyellow-jp', 'pokeyellow', 'pokeyellowCHS-sjp', 'yellow', 'YEJP')]:
        game = 'red' if jp == 'pokegreen' else ('blue' if jp == 'pokeblue' else 'yellow')
        japanese, _ = extract_assembly(jp, game)
        english, edges = extract_assembly(en, game)
        chinese, replacements = extract_workbooks(zh, ver, game)
        byenglish, bylabel, enlabels = collections.defaultdict(list), collections.defaultdict(list), collections.defaultdict(list)
        for b in chinese:
            byenglish[direct_key(b['english'])].append(b)
            bylabel[label_key(b['label'])].append(b)
        for e in english:
            enlabels[e['label']].append(e)
        counts = collections.Counter()
        for j in japanese:
            if not j['file'].startswith(('scripts/', 'engine/', 'home/')):
                continue
            ja, je, jv, ju = render(j['ops'], 'ja')
            if not ja.strip():
                continue
            rid = f"{jp}:{j['file']}:{j['label']}"
            same_label = sum(r['id'].split('#')[0] == rid for r in all_records)
            if same_label:
                rid += '#'+str(same_label+1)
            names = edges.get(j['label'], []) or [j['label'], '_'+j['label']]
            candidate_groups, evidence = [], []
            for name in names:
                group = {}
                for e in enlabels.get(name, []):
                    ek = english_key(quoted_text(e['ops']))
                    bs = byenglish.get(ek, []) if ek else []
                    exact_labels = [b for b in bs if label_key(b['label']) == label_key(e['label'])]
                    same_file = [b for b in bs if b['file'] == e['file']]
                    for b in exact_labels or same_file or bs:
                        bid = f"{b['book']}:{b['sheet']}:{b['row']}"
                        group[bid] = b
                        evidence.append(dict(english_file=e['file'], english_label=e['label'],
                                             english_line=e['line'], bridge='text_far+english_text'))
                if group:
                    candidate_groups.append(group)
            candidates = candidate_groups[0] if len(candidate_groups) == 1 else {}
            method = 'upstream_pointer_and_english_text'
            for override in overrides:
                if override['japanese_label'] == j['label'] and override['japanese_repo'] in {'*', jp}:
                    candidate_groups = [{f"{b['book']}:{b['sheet']}:{b['row']}": b
                                         for b in bylabel[legacy]} for legacy in override['chinese_label'].split('|')]
                    method = override['method']
                    evidence = [dict(note=override['evidence'], reviewer='model')]
            if not candidate_groups:
                candidates = {}
                for name in names:
                    for b in bylabel.get(label_key(name), []):
                        candidates[f"{b['book']}:{b['sheet']}:{b['row']}"] = b
                method = 'label_candidate'
                candidate_groups = [candidates]
            record = dict(id=rid, games=games, kind='dialogue', japanese_repo=jp,
                          japanese_file=j['file'], japanese_label=j['label'], japanese_line=j['line'],
                          japanese=ja, japanese_commands=packed(j['ops']), japanese_events=packed(je),
                          japanese_termination=j['termination'], chinese_repo=zh,
                          chinese_source='', chinese_label='', chinese='', chinese_cells='',
                          chinese_events='', variables=packed(jv), method=method,
                          evidence=packed(evidence), mapping_status='unresolved',
                          segment_status='unavailable', variable_status='unresolved',
                          render_status='unsupported' if ju else 'parsed')
            if candidate_groups and all(len(group) == 1 for group in candidate_groups):
                selected = [next(iter(group.values())) for group in candidate_groups]
                zt, ze, zv, zu = render([o for b in selected for o in b['ops']], 'zh', quoted=False)
                record.update(chinese_source=packed([f"src/xlsx/{b['book']}#{b['sheet']}!F{b['row']}:G{max(o[2] for o in b['ops'])}" for b in selected]),
                              chinese_label=';'.join(b['label'] for b in selected), chinese=zt,
                              chinese_cells=packed([dict(book=b['book'],sheet=b['sheet'],cells=b['cells']) for b in selected]),
                              chinese_events=packed(ze), mapping_status='aligned' if method != 'label_candidate' else 'review',
                              render_status='unsupported' if ju or zu else 'parsed')
                # Slots are positional only after the corresponding full source
                # texts are known; multiple mutable buffers need semantic review.
                jtypes = [v['kind'] for v in jv if v['kind'] != 'NAME']
                ztypes = [v['kind'] for v in zv if v['kind'] != 'NAME']
                record['variable_status'] = 'none' if not jv and not zv else 'review'
                if jtypes == ztypes and len(jtypes) <= 1:
                    jnames = {v['slot'] for v in jv if v['kind'] == 'NAME'}
                    znames = {v['slot'] for v in zv if v['kind'] == 'NAME'}
                    if jnames == znames:
                        record['variable_status'] = 'typed_unique'
                if j['label'] in bindings:
                    binding = bindings[j['label']]
                    record['variable_status'] = 'model_semantic'
                    record['evidence'] = packed(dict(alignment=evidence, variable_binding=binding))
                    if binding['rule'] == 'swap_ram1_ram2':
                        record['chinese'] = record['chinese'].replace('{{RAM1}}','{{TEMP}}').replace('{{RAM2}}','{{RAM1}}').replace('{{TEMP}}','{{RAM2}}')
                        for v in zv:
                            v['slot'] = {'RAM1':'RAM2','RAM2':'RAM1'}.get(v['slot'],v['slot'])
                    elif binding['rule'] == 'require_PLAYER':
                        record['variable_status'] = 'requires_observed_PLAYER'
                for language, vs in [('ja', jv), ('zh', zv)]:
                    for v in vs:
                        all_variables.append(dict(record_id=rid, language=language, **v))
                if record['mapping_status'] == 'aligned' and record['render_status'] == 'parsed' and j['termination'] != 'asm':
                    # Full source mapping and visible segmentation are separate.
                    # Only a whole first screen is intrinsically safe here.
                    if len(ja.splitlines()) <= 2 and '\n\n' not in ja:
                        record['segment_status'] = 'whole_visible_unit'
            else:
                record['mapping_status'] = 'ambiguous' if any(candidate_groups) else 'unresolved'
            if j['label'] == 'UnknownText_2812':
                record.update(mapping_status='source_unreferenced',kind='unreferenced',
                              evidence='home/text_script.asm explicitly marks this label unreferenced; no runtime entry reference')
            if record['mapping_status'] != 'aligned' or record['render_status'] != 'parsed' or record['variable_status'] == 'review':
                issues.append(dict(record_id=rid, reason=record['mapping_status']+';'+record['render_status']+';'+record['variable_status'],
                                   japanese=ja, candidates=packed([[dict(label=b['label'],book=b['book'],sheet=b['sheet'],row=b['row']) for b in group.values()] for group in candidate_groups])))
            all_records.append(record)
            counts[record['mapping_status']] += 1
            # Reconstruct the two-row dialogue window. @ never adds a line;
            # PARA clears it, CONT scrolls it. Chinese slicing is not inferred.
            lines = ja.split('\n')
            visible, seq = [], 0
            for line in lines:
                if line == '':
                    visible = []
                    continue
                visible = (visible+[line])[-2:]
                seq += 1
                all_fragments.append(dict(record_id=rid, ordinal=seq, japanese='\n'.join(visible),
                                          previous_ordinal=seq-1 if seq>1 else '', chinese='',
                                          chinese_start='', chinese_end='', status='unavailable',
                                          source='game_text_controls'))
            if record['segment_status'] == 'whole_visible_unit' and seq:
                f = all_fragments[-1]
                f.update(chinese=record['chinese'], chinese_start=0, chinese_end=len(record['chinese']),
                         status='whole_visible_unit')
        stats[jp] = dict(counts)
        add_words(jp, zh, games, replacements, all_records, all_fragments)
        add_fixed_words(jp, zh, games, replacements, all_records, all_fragments)
        add_compositions(jp, all_records, all_fragments)
    # Japanese Yellow retains the evolution statement. Its own CHS workbook
    # uses a general compliment here; the selected RGB source has the precise
    # existing evolution sentence. Reuse it verbatim, without rewriting Chinese.
    blue_trade=next(r for r in all_records if r['japanese_repo']=='pokeblue' and r['japanese_label']=='AfterTrade2Text')
    yellow_trade=next(r for r in all_records if r['japanese_repo']=='pokeyellow-jp' and r['japanese_label']=='AfterTrade2Text')
    for field in ['chinese_repo','chinese_source','chinese_label','chinese','chinese_cells','chinese_events']:
        yellow_trade[field]=blue_trade[field]
    yellow_trade.update(method='model_semantic_cross_source',
                        evidence='Japanese Yellow says the traded Pokemon evolved; reuse the unchanged RGB evolution sentence rather than the Yellow generic compliment')
    for fragment in all_fragments:
        if fragment['record_id']==yellow_trade['id'] and fragment['status']=='whole_visible_unit':
            fragment.update(chinese=yellow_trade['chinese'],chinese_end=len(yellow_trade['chinese']))
    issues = [dict(record_id=r['id'], reason=r['mapping_status']+';'+r['render_status']+';'+r['variable_status'],
                   japanese=r['japanese'], candidates=r['chinese_label']) for r in all_records
              if r['mapping_status'] not in {'aligned','component','source_unreferenced'} or r['render_status']!='parsed'
              or (r['variable_status']=='review' and r['mapping_status']!='component')]
    component_ids={r['id'] for r in all_records if r['mapping_status']=='component'}
    for fragment in all_fragments:
        if fragment['record_id'] in component_ids:
            fragment.update(status='use_composition',chinese='',chinese_start='',chinese_end='')
    stats = {jp: dict(collections.Counter(r['mapping_status'] for r in all_records if r['japanese_repo']==jp))
             for jp in ['pokegreen','pokeblue','pokeyellow-jp']}
    for r in all_records:
        r['japanese_ocr'] = r['japanese'].replace('　',' ').replace('⋯','…').translate(str.maketrans('０１２３４５６７８９','0123456789'))
        r['japanese_ocr'] = '\n'.join(line.rstrip() for line in r['japanese_ocr'].split('\n'))
    for f in all_fragments:
        f['japanese_ocr'] = f['japanese'].replace('　',' ').replace('⋯','…').translate(str.maketrans('０１２３４５６７８９','0123456789'))
        f['japanese_ocr'] = '\n'.join(line.rstrip() for line in f['japanese_ocr'].split('\n'))
    write_csv(DEST/'sources.csv', source_rows, list(source_rows[0]))
    write_csv(DEST/'records.csv', all_records, list(all_records[0]))
    write_csv(DEST/'fragments.csv', all_fragments, list(all_fragments[0]))
    write_csv(DEST/'variables.csv', all_variables, ['record_id','language','slot','kind','address','format','offset','source_line'])
    write_csv(ROOT/'reports/text-corpus-issues.csv', issues, ['record_id','reason','japanese','candidates'])
    report=dict(sources={r['id']:r['commit'] for r in source_rows},records=len(all_records),
                mapping_status=stats,kinds=dict(collections.Counter(r['kind'] for r in all_records)),
                fragments=dict(collections.Counter(f['status'] for f in all_fragments)),
                variable_status=dict(collections.Counter(r['variable_status'] for r in all_records)),
                unresolved_source_alignment_issues=len(issues),verification='static_source_and_data_construction',
                dynamic_tests_run=False,runtime_translation_integrated=False,
                chinese_acceptance_passed=False)
    (ROOT/'reports/text-corpus.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (ROOT/'output/text-corpus-summary.json').write_text(packed(stats)+'\n', encoding='utf-8')
    print(json.dumps(stats, ensure_ascii=False, indent=2))
    print('records',len(all_records),'fragments',len(all_fragments),'issues',len(issues))


if __name__ == '__main__':
    build()
